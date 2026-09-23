# -*- coding: utf-8 -*-
"""
敏實科技大學 USR 計畫 - 活動記錄表自動生成器 (CLI / 模組)
Activity Record Generator for MINTH University USR Projects
"""

import os
import sys
import json
import argparse
import datetime
from PIL import Image
import docx
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def get_minguo_timestamp():
    """產出民國年11碼時間戳記：[民國年3碼][月2碼][日2碼][時2碼][分2碼]"""
    now = datetime.datetime.now()
    m_year = now.year - 1911
    return f"{m_year:03d}{now.month:02d}{now.day:02d}{now.hour:02d}{now.minute:02d}"

def apply_font(run, font_name="標楷體", font_size_pt=12, bold=False, color_rgb=None):
    """統一套用標楷體與英數字型規範"""
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
    run._element.rPr.rFonts.set(qn('w:ascii'), 'Times New Roman')
    run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Times New Roman')
    if font_size_pt:
        run.font.size = Pt(font_size_pt)
    run.bold = bold
    if color_rgb:
        run.font.color.rgb = color_rgb

def set_cell_text(cell, text, font_size_pt=12, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, line_spacing=1.15, space_after=0):
    """設置儲存格單行/基本文字"""
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    if line_spacing:
        p.paragraph_format.line_spacing = line_spacing
    if space_after:
        p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    apply_font(run, font_name="標楷體", font_size_pt=font_size_pt, bold=bold)
    return p

def set_cell_multiline(cell, lines, font_size_pt=10.5, line_spacing=1.2, space_after=2):
    """設置儲存格多行文字段落"""
    cell.text = ""
    for i, line in enumerate(lines):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.paragraph_format.line_spacing = line_spacing
        p.paragraph_format.space_after = Pt(space_after)
        run = p.add_run(line)
        apply_font(run, font_name="標楷體", font_size_pt=font_size_pt)

def set_cell_image(cell, image_path, max_width_cm=8.3, max_height_cm=5.4):
    """將照片等比例縮放居中嵌入儲存格"""
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    
    if not os.path.exists(image_path):
        run = p.add_run(f"[照片未找到: {os.path.basename(image_path)}]")
        apply_font(run, font_size_pt=10, bold=True)
        return

    try:
        with Image.open(image_path) as im:
            w_px, h_px = im.size
            aspect = w_px / h_px
            target_aspect = max_width_cm / max_height_cm
            if aspect >= target_aspect:
                final_w = Cm(max_width_cm)
                final_h = Cm(max_width_cm / aspect)
            else:
                final_h = Cm(max_height_cm)
                final_w = Cm(max_height_cm * aspect)
        run = p.add_run()
        run.add_picture(image_path, width=final_w, height=final_h)
    except Exception as e:
        run = p.add_run(f"[照片載入失敗: {e}]")
        apply_font(run, font_size_pt=10, bold=True)

def generate_activity_record(data, template_path=None, output_path=None):
    """
    主要活動記錄表生成器
    :param data: 包含所有填寫資訊的字典 (dict)
    :param template_path: 範本 docx 檔案路徑 (若為 None 則自動尋找預設範本)
    :param output_path: 產出 docx 檔案路徑 (若為 None 則自動依全域規範命名存入 Spark 目錄)
    :return: 產出的 docx 檔案完整路徑
    """
    if not template_path:
        base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        template_path = os.path.join(base_dir, "templates", "template_activity_record.docx")
        
    if not os.path.exists(template_path):
        # Fallback to local draft template
        fallback = r"D:\AI_Python_草稿\template_activity_record.docx"
        if os.path.exists(fallback):
            template_path = fallback
        else:
            raise FileNotFoundError(f"找不到活動記錄表標準範本: {template_path}")

    doc = docx.Document(template_path)
    
    # 1. 更新大標題年度
    year_str = str(data.get("year", "115"))
    for p in doc.paragraphs:
        if "社會責任實踐計畫」活動記錄表" in p.text:
            p.text = f"「{year_str}年度社會責任實踐計畫」活動記錄表"
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                apply_font(r, font_name="標楷體", font_size_pt=20, bold=True)
            break

    # 2. Table 0：活動基本資訊與紀要
    t0 = doc.tables[0]
    
    # R0: USR子計畫勾選
    subproj_id = str(data.get("subproject_id", "1"))
    sp1 = "■ 子計畫1：創生共學AI及ESG智慧創新平台" if subproj_id == "1" else "□ 子計畫1：創生共學AI及ESG智慧創新平台"
    sp4 = "■ 子計畫4：AI智慧減碳旅遊城鄉體驗" if subproj_id == "4" else "□ 子計畫4：AI智慧減碳旅遊城鄉體驗"
    sp2 = "■ 子計畫2：AI技術應用在農業領域合作" if subproj_id == "2" else "□ 子計畫2：AI技術應用在農業領域合作"
    sp5 = "■ 子計畫5：AI技術應用在碳中和技術領域合作" if subproj_id == "5" else "□ 子計畫5：AI技術應用在碳中和技術領域合作"
    sp3 = "■ 子計畫3：AI及ESG食農教育特色餐飲" if subproj_id == "3" else "□ 子計畫3：AI及ESG食農教育特色餐飲"
    sph = "■ 校務Hub：" if subproj_id.lower() in ["hub", "校務hub"] else "□ 校務Hub："
    
    lines_subproj = [
        f"{sp1}    {sp4}",
        f"{sp2}    {sp5}",
        f"{sp3}    {sph}"
    ]
    set_cell_multiline(t0.rows[0].cells[1], lines_subproj, font_size_pt=9.5, line_spacing=1.15, space_after=1)

    # R1: 子計畫主持人 / 工讀生
    host_name = data.get("host_name", "")
    assistant_name = data.get("assistant_name", "")
    set_cell_text(t0.rows[1].cells[1], host_name, font_size_pt=12, align=WD_ALIGN_PARAGRAPH.LEFT)
    set_cell_text(t0.rows[1].cells[3], assistant_name, font_size_pt=12, align=WD_ALIGN_PARAGRAPH.LEFT)

    # R2: 活動名稱
    act_name = data.get("activity_name", "")
    set_cell_text(t0.rows[2].cells[0], f"活動名稱：{act_name}", font_size_pt=12, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT)

    # R3: 日期/時間 與 地點
    date_time = data.get("date_time", "")
    location = data.get("location", "")
    set_cell_text(t0.rows[3].cells[0], f"日期/時間：{date_time}", font_size_pt=11, align=WD_ALIGN_PARAGRAPH.LEFT)
    set_cell_text(t0.rows[3].cells[2], f"地點：{location}", font_size_pt=11, align=WD_ALIGN_PARAGRAPH.LEFT)

    # R4: 主講者/與會人
    speaker_attendees = data.get("speaker_attendees", "")
    set_cell_text(t0.rows[4].cells[0], f"主講者/與會人：{speaker_attendees}", font_size_pt=11, align=WD_ALIGN_PARAGRAPH.LEFT)

    # R5: 執行成效 (量化指標 + 質化指標)
    quant = data.get("quantitative", {})
    t_in = quant.get("inner_teacher", 0)
    s_in = quant.get("inner_student", 0)
    a_in = quant.get("inner_assistant", 0)
    t_out = quant.get("cross_teacher", 0)
    s_out = quant.get("cross_student", 0)
    a_out = quant.get("cross_assistant", 0)
    f_part = quant.get("field_partner", 0)
    c_gov = quant.get("central_gov", 0)
    l_gov = quant.get("local_gov", 0)
    ngo = quant.get("ngo_npo", 0)
    q_sent = quant.get("survey_sent", 0)
    q_valid = quant.get("survey_valid", 0)

    qual_text = data.get("qualitative_summary", "")

    perf_lines = [
        "量化指標：",
        f"  校內參與：教師：{t_in} 人、學生：{s_in} 人、專任助理：{a_in} 人",
        f"  跨校參與：教師：{t_out} 人、學生：{s_out} 人、專任助理：{a_out} 人",
        f"  場域參與：場域夥伴人數：{f_part} 人",
        f"  其他參與：中央政府人數：{c_gov} 人、地方政府人數：{l_gov} 人、NGO/NPO：{ngo} 人",
        "",
        "滿意度調查：",
        f"  發放問卷數：{q_sent} 份、有效問卷數：{q_valid} 份",
        "",
        "質化指標：",
        qual_text
    ]
    set_cell_multiline(t0.rows[5].cells[1], perf_lines, font_size_pt=10.5, line_spacing=1.15, space_after=1)

    # R6: 活動紀要
    summary_content = data.get("activity_summary", "")
    if isinstance(summary_content, str):
        summary_lines = [line.strip() for line in summary_content.split("\n") if line.strip()]
    else:
        summary_lines = summary_content
    set_cell_multiline(t0.rows[6].cells[1], summary_lines, font_size_pt=10.5, line_spacing=1.2, space_after=2)

    # R7: 檢討與建議
    review_content = data.get("review_suggestions", "")
    if isinstance(review_content, str):
        review_lines = [line.strip() for line in review_content.split("\n") if line.strip()]
    else:
        review_lines = review_content
    set_cell_multiline(t0.rows[7].cells[1], review_lines, font_size_pt=10.5, line_spacing=1.2, space_after=2)

    # 3. Table 1：活動剪影照片矩陣與相關附件
    t1 = doc.tables[1]
    photos = data.get("photos", [])  # [{"path": "...", "caption": "..."}]
    num_photos = min(len(photos), 6)
    row_pairs = [(0, 1), (2, 3), (4, 5)]

    for idx in range(6):
        pair_idx = idx // 2
        col_idx = 0 if (idx % 2 == 0) else 2
        img_row_idx, cap_row_idx = row_pairs[pair_idx]
        
        if idx < num_photos:
            p_info = photos[idx]
            p_path = p_info.get("path", "")
            p_cap = p_info.get("caption", "")
            if not p_cap.startswith("說明："):
                p_cap = f"說明：{p_cap}"
                
            set_cell_image(t1.rows[img_row_idx].cells[col_idx], p_path)
            set_cell_text(t1.rows[cap_row_idx].cells[col_idx], p_cap, font_size_pt=10.5, align=WD_ALIGN_PARAGRAPH.LEFT)
        else:
            t1.rows[img_row_idx].cells[col_idx].text = ""
            t1.rows[cap_row_idx].cells[col_idx].text = ""

    # 若照片僅有 2 張或 4 張，自後往前移除多餘空白列
    if num_photos <= 2:
        for r_del in [5, 4, 3, 2]:
            tr = t1.rows[r_del]._tr
            tr.getparent().remove(tr)
    elif num_photos <= 4:
        for r_del in [5, 4]:
            tr = t1.rows[r_del]._tr
            tr.getparent().remove(tr)

    # 附件核選列（動態對齊為 Table 1 最後一列）
    last_row = t1.rows[-1]
    attachments = data.get("attachments", ["簽到單", "宣傳海報"])
    att_agenda = "■ 議程" if "議程" in attachments else "□ 議程"
    att_signin = "■ 簽到單" if "簽到單" in attachments else "□ 簽到單"
    att_poster = "■ 宣傳海報" if "宣傳海報" in attachments else "□ 宣傳海報"
    att_slides = "■ 簡報" if "簡報" in attachments else "□ 簡報"
    att_other = "□ 其他___________"
    for a in attachments:
        if a not in ["議程", "簽到單", "宣傳海報", "簡報"]:
            att_other = f"■ 其他：{a}"
            break
            
    att_text = f"{att_agenda}   {att_signin}   {att_poster}   {att_slides}   {att_other}   (請勾選可提供之附件)"
    set_cell_text(last_row.cells[1], att_text, font_size_pt=10.5, align=WD_ALIGN_PARAGRAPH.LEFT)

    # 4. 擴充全頁佐證附件 (如簽到表掃描檔、宣傳海報)
    annex_list = data.get("annex_files", [])  # [{"title": "簽到表(掃描檔)", "path": "..."}]
    for annex in annex_list:
        doc.add_page_break()
        title_p = doc.add_paragraph()
        title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_p.paragraph_format.space_before = Pt(6)
        title_p.paragraph_format.space_after = Pt(12)
        r = title_p.add_run(annex.get("title", "活動附件"))
        apply_font(r, font_name="標楷體", font_size_pt=14, bold=True)
        
        a_path = annex.get("path", "")
        if os.path.exists(a_path):
            img_p = doc.add_paragraph()
            img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            try:
                img_p.add_run().add_picture(a_path, width=Cm(16.0))
            except Exception as e:
                img_p.add_run(f"[附件載入失敗: {e}]")

    # 5. 檔名與存檔路徑 (嚴格依全域規範：民國年月日時分 11 碼 + -anti.docx)
    if not output_path:
        out_dir = r"E:\我的雲端硬碟\Gemini_Spark_產出檔案"
        if not os.path.exists(out_dir):
            os.makedirs(out_dir, exist_ok=True)
        ts = get_minguo_timestamp()
        safe_act_name = "".join(c for c in act_name if c.isalnum() or c in ['_', '-'])[:20]
        filename = f"USR活動記錄表_子計畫{subproj_id}_{safe_act_name}-{ts}-anti.docx"
        output_path = os.path.join(out_dir, filename)

    doc.save(output_path)
    return output_path

def main():
    parser = argparse.ArgumentParser(description="敏實科技大學 USR 計畫活動記錄表生成器")
    parser.add_argument("--data", "-d", help="活動資料 JSON 檔案路徑", required=True)
    parser.add_argument("--template", "-t", help="範本 docx 檔案路徑 (選填)")
    parser.add_argument("--output", "-o", help="產出 docx 檔案路徑 (選填)")
    args = parser.parse_args()

    with open(args.data, "r", encoding="utf-8") as f:
        data = json.load(f)

    out_file = generate_activity_record(data, template_path=args.template, output_path=args.output)
    print(f"[OK] 活動記錄表成功生成：{out_file}")

if __name__ == '__main__':
    main()
